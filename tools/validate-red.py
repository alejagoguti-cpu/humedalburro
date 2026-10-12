"""Valida la calidad de la red: grados, imagenes, coordenadas, tensiones."""
import openpyxl, os, sys

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'datos/red-pot-kennedy-visor.xlsx'
wb = openpyxl.load_workbook(XLSX, data_only=True)

def txt(v):
    return '' if v is None else str(v).strip()

def rows(ws, fc, lc):
    out = []
    for r in range(2, ws.max_row + 1):
        row = [txt(ws.cell(r, c).value) for c in range(fc, lc + 1)]
        if row[0]:
            out.append((r, row))
    return out

nodos = rows(wb['NODOS'], 1, 18)
aristas = rows(wb['ARISTAS'], 1, 12)
ids = [r[0] for _, r in nodos]
print('NODOS:', len(nodos), 'ARISTAS:', len(aristas))

# duplicados
print('IDs duplicados:', len(ids) - len(set(ids)))

# grados
from collections import Counter
deg = Counter()
for _, r in aristas:
    deg[r[1]] += 1
    deg[r[2]] += 1
aislados = [i for i in ids if deg[i] == 0]
print('Nodos aislados (grado 0):', aislados)

# autoaristas
auto = [r[0] for _, r in aristas if r[1] == r[2]]
print('Autoaristas:', auto)

# tensiones validas
tens = Counter(r[6] for _, r in aristas)
print('Tensiones:', dict(tens))

# pesos
pesos = Counter(r[7] for _, r in aristas)
print('Pesos:', dict(pesos))

# tipos usados
tipos = Counter(r[3] for _, r in aristas)
print('Tipos usados:', dict(tipos))

# capas usadas
capas = Counter(r[2] for _, r in nodos)
print('Capas usadas:', dict(capas))

# imagenes
falt_img_nodo = []
for ln, r in nodos:
    for col, key in ((12, 'img'), (15, 'img1'), (16, 'img2')):
        p = r[col - 1]
        if p and not os.path.exists(p):
            falt_img_nodo.append((r[0], key, p))
print('Imagenes de nodos faltantes:', len(falt_img_nodo))
for x in falt_img_nodo[:30]:
    print('  ', x)

# coordenadas
fuera = []
for ln, r in nodos:
    try:
        lat = float(r[12]); lon = float(r[13])
    except ValueError:
        print('  COORD INVALIDA fila', ln, r[0], repr(r[12]), repr(r[13]))
        continue
    if not (4.5 <= lat <= 4.75 and -74.25 <= lon <= -74.05):
        fuera.append((r[0], lat, lon))
print('Coords fuera de Kennedy aprox:', fuera)

# campos vacios criticos
for ln, r in nodos:
    for ci, nm in ((2, 'nombre'), (3, 'capa'), (7, 'role'), (8, 'alert'), (11, 'fuente')):
        if not r[ci - 1]:
            print('  NODO VACIO', r[0], nm, 'fila', ln)
for ln, r in aristas:
    for ci, nm in ((4, 'tipo'), (6, 'rationale')):
        if not r[ci - 1]:
            print('  ARISTA VACIA', r[0], nm, 'fila', ln)
print('VALIDACION OK')
