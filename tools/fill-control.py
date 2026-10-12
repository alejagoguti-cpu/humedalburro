"""Llena la hoja CONTROL de datos/red-pot-kennedy-visor.xlsx con la validacion actual."""
from collections import Counter
from datetime import date

import openpyxl

XLSX = 'datos/red-pot-kennedy-visor.xlsx'
wb = openpyxl.load_workbook(XLSX)


def txt(v):
    return '' if v is None else str(v).strip()


def rows(ws, fc, lc):
    out = []
    for r in range(2, ws.max_row + 1):
        row = [txt(ws.cell(r, c).value) for c in range(fc, lc + 1)]
        if row[0]:
            out.append(row)
    return out


capas = {r[1] for r in rows(wb['CAPAS'], 1, 6)}
tipos = {r[1] for r in rows(wb['INTERACCIONES'], 1, 5)}
nodos = rows(wb['NODOS'], 1, 18)
aristas = rows(wb['ARISTAS'], 1, 12)
ids = [r[0] for r in nodos]
idset = set(ids)

checks = [
    ('Nodos con ID duplicado', len(ids) - len(idset)),
    ('Nodos con capa no valida', sum(1 for r in nodos if r[2] not in capas)),
    ('Aristas con origen inexistente', sum(1 for r in aristas if r[1] not in idset)),
    ('Aristas con destino inexistente', sum(1 for r in aristas if r[2] not in idset)),
    ('Aristas con tipo fuera de catalogo', sum(1 for r in aristas if r[3] not in tipos)),
    ('Autoaristas (origen = destino)', sum(1 for r in aristas if r[1] == r[2])),
]
deg = Counter()
for r in aristas:
    deg[r[1]] += 1
    deg[r[2]] += 1
checks.append(('Nodos aislados (grado 0)', sum(1 for i in ids if deg[i] == 0)))

ws = wb['CONTROL']
by_name = {}
for r in range(2, ws.max_row + 1):
    name = txt(ws.cell(r, 1).value)
    if name:
        by_name[name] = r

for name, count in checks:
    if name in by_name:
        ws.cell(by_name[name], 2).value = count
        ws.cell(by_name[name], 3).value = 'OK' if count == 0 else 'REVISAR'

ok = all(c == 0 for _, c in checks)
if 'ESTADO GLOBAL' in by_name:
    ws.cell(by_name['ESTADO GLOBAL'], 2).value = 'OK' if ok else 'CON ERRORES'
    ws.cell(by_name['ESTADO GLOBAL'], 3).value = str(date.today())
if 'Total nodos' in by_name:
    ws.cell(by_name['Total nodos'], 2).value = len(nodos)
if 'Total aristas' in by_name:
    ws.cell(by_name['Total aristas'], 2).value = len(aristas)

wb.save(XLSX)
print('CONTROL actualizado:', 'OK' if ok else 'CON ERRORES',
      '| nodos:', len(nodos), '| aristas:', len(aristas))
for name, count in checks:
    print(f'  {name}: {count}')
