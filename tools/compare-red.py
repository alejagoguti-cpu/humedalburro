import openpyxl, re, json, sys

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'datos/red-pot-kennedy-visor.xlsx'
HTML = sys.argv[2] if len(sys.argv) > 2 else 'index.html'
wb = openpyxl.load_workbook(XLSX, data_only=True)

def txt(v):
    return '' if v is None else str(v).strip()

def rows(ws, fc, lc):
    out = []
    for r in range(2, ws.max_row + 1):
        row = [txt(ws.cell(r, c).value) for c in range(fc, lc + 1)]
        if row[0]:
            out.append(row)
    return out

capas = rows(wb['CAPAS'], 1, 6)
tipos = rows(wb['INTERACCIONES'], 1, 5)
nodos = rows(wb['NODOS'], 1, 18)
aristas = rows(wb['ARISTAS'], 1, 12)

html = open(HTML, encoding='utf-8').read()
a = '<!--RED_DATA_START-->'
b = '<!--RED_DATA_END-->'
seg = html[html.index(a):html.index(b)]
d = json.loads(re.search(r'window\.RED_DATA\s*=\s*(\{.*?\});', seg, re.S).group(1))

print('XLSX : capas', len(capas), 'tipos', len(tipos), 'nodos', len(nodos), 'aristas', len(aristas))
print('EMB  : capas', len(d['categories']), 'tipos', len(d['interactions']), 'nodos', len(d['nodes']), 'aristas', len(d['edges']))

capaIdx = {c[1]: i for i, c in enumerate(capas)}
tipoIdx = {t[1]: i for i, t in enumerate(tipos)}

def xnode(r):
    icon = r[17] if r[17] else (capas[capaIdx[r[2]]][4] if r[2] in capaIdx else '')
    return {'id': r[0], 'name': r[1], 'cat': capaIdx.get(r[2]), 'sciname': r[3],
            'scale': r[4], 'loc': r[5], 'role': r[6], 'alert': r[7], 'actors': r[8],
            'hypothesis': r[9], 'source': r[10], 'img': r[11], 'img1': r[14],
            'img2': r[15], 'icon': icon}

dif = 0
for xr, en in zip(nodos, d['nodes']):
    xn = xnode(xr)
    for k, v in xn.items():
        if str(en.get(k, '')) != str(v):
            print('NODO DIFF', xr[0], k, repr(v)[:70], 'vs', repr(en.get(k, ''))[:70])
            dif += 1
            break
    if str(en['geo']['lat']) != str(xr[12]) or str(en['geo']['lon']) != str(xr[13]):
        print('GEO DIFF', xr[0], xr[12], xr[13], 'vs', en['geo'])
        dif += 1
print('nodos con diff:', dif)

idIdx = {n['id']: i for i, n in enumerate(d['nodes'])}
edif = 0
for xr, ee in zip(aristas, d['edges']):
    if idIdx.get(xr[1]) != ee['source'] or idIdx.get(xr[2]) != ee['target'] or tipoIdx.get(xr[3]) != ee['type']:
        print('ARISTA DIFF', xr[0], xr[1], xr[2], xr[3])
        edif += 1
    for k, xk in [('detail', 4), ('rationale', 5), ('tension', 6), ('declared', 8),
                  ('operative', 9), ('hypothesis', 10), ('evidence', 11)]:
        if str(ee.get(k, '')) != str(xr[xk]):
            print('ARISTA FIELD DIFF', xr[0], k)
            edif += 1
            break
print('aristas con diff:', edif)

kx = sorted(set(r[0] for r in rows(wb['RED'], 1, 2) if r[0]))
print('config xlsx n=', len(kx))
print('config emb keys:', sorted(d['config'].keys()))
