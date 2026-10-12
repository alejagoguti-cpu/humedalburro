#!/usr/bin/env python3
# Convierte datos/red-pot-kennedy-visor.xlsx -> embed en la pagina HTML.
# Uso:  python tools/build-red.py [libro.xlsx] [pagina.html]
# Por defecto: datos/red-pot-kennedy-visor.xlsx -> index.html. Los datos quedan EMBEBIDOS
# en la pagina (entre <!--RED_DATA_START--> y <!--RED_DATA_END-->).
# Port fiel de tools/build-red.js (misma matematica determinista: Brandes, Louvain,
# layout 3D Fruchterman-Reingold con semilla fija y proyeccion WGS84 -> UTM 18N -> escena).
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'datos' / 'red-pot-kennedy-visor.xlsx'
HTML_PATH = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'index.html'


def txt(v):
    if v is None:
        return ''
    return str(v).strip()


def rows(ws, first_col, last_col):
    out = []
    for r in range(2, ws.max_row + 1):
        row = [txt(ws.cell(r, c).value) for c in range(first_col, last_col + 1)]
        if row[0]:
            out.append(row)
    return out


def to_fixed(v, digits):
    """Replica Number.prototype.toFixed (mitad lejos de cero)."""
    f = 10 ** digits
    if v >= 0:
        return math.floor(v * f + 0.5) / f
    return math.ceil(v * f - 0.5) / f


def jsnum(v):
    """Replica JSON.stringify de JS: los enteros se imprimen sin decimales."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def make_rng(seed):
    s = seed & 0xFFFFFFFF

    def rand():
        nonlocal s
        s = (s * 1664525 + 1013904223) & 0xFFFFFFFF
        return s / 4294967296

    return rand


def brandes(n, adj):
    bc = [0.0] * n
    for s in range(n):
        st = []
        pred = [[] for _ in range(n)]
        sigma = [0] * n
        dist = [-1] * n
        sigma[s] = 1
        dist[s] = 0
        q = [s]
        while q:
            v = q.pop(0)
            st.append(v)
            for w in adj[v]:
                if dist[w] < 0:
                    dist[w] = dist[v] + 1
                    q.append(w)
                if dist[w] == dist[v] + 1:
                    sigma[w] += sigma[v]
                    pred[w].append(v)
        delta = [0.0] * n
        while st:
            w = st.pop()
            for v in pred[w]:
                delta[v] += (sigma[v] / sigma[w]) * (1 + delta[w])
            if w != s:
                bc[w] += delta[w]
    mx = max(bc + [1e-9])
    return [jsnum(to_fixed(b / mx, 4)) for b in bc]


def louvain(n, edges):
    g_n = n
    g_w = [dict() for _ in range(n)]
    for a, b, w in edges:
        g_w[a][b] = g_w[a].get(b, 0) + w
        g_w[b][a] = g_w[b].get(a, 0) + w
    m2 = sum(2 * e[2] for e in edges)
    mapping = list(range(n))
    for _level in range(10):
        k = [sum(m.values()) for m in g_w]
        comm = list(range(g_n))
        tot = k[:]
        moved = True
        guard = 0
        while moved and guard < 50:
            guard += 1
            moved = False
            for i in range(g_n):
                ci = comm[i]
                links = {}
                for j, w in g_w[i].items():
                    if j != i:
                        links[comm[j]] = links.get(comm[j], 0) + w
                tot[ci] -= k[i]
                best = ci
                best_gain = links.get(ci, 0) - tot[ci] * k[i] / m2
                for c, w in links.items():
                    gain = w - tot[c] * k[i] / m2
                    if gain > best_gain + 1e-12:
                        best = c
                        best_gain = gain
                tot[best] += k[i]
                if best != ci:
                    comm[i] = best
                    moved = True
        ids = list(dict.fromkeys(comm))
        remap = {c: i for i, c in enumerate(ids)}
        if len(ids) == g_n:
            break
        mapping = [remap[comm[x]] for x in mapping]
        ng_w = [dict() for _ in range(len(ids))]
        for i in range(g_n):
            for j, w in g_w[i].items():
                a = remap[comm[i]]
                b = remap[comm[j]]
                ng_w[a][b] = ng_w[a].get(b, 0) + w
        g_n = len(ids)
        g_w = ng_w
    deg = [0] * n
    for a, b, w in edges:
        deg[a] += w
        deg[b] += w
    Q = 0.0
    sum_in = {}
    sum_tot = {}
    for a, b, w in edges:
        if mapping[a] == mapping[b]:
            sum_in[mapping[a]] = sum_in.get(mapping[a], 0) + 2 * w
    for i, dval in enumerate(deg):
        sum_tot[mapping[i]] = sum_tot.get(mapping[i], 0) + dval
    for c in sum_tot:
        Q += (sum_in.get(c, 0)) / m2 - (sum_tot[c] / m2) ** 2
    size = {}
    for c in mapping:
        size[c] = size.get(c, 0) + 1
    order = sorted(sorted(size.keys()), key=lambda c: size[c], reverse=True)
    fin = {c: i for i, c in enumerate(order)}
    return [fin[c] for c in mapping], jsnum(to_fixed(Q, 2))


def layout3d(n, edges):
    rand = make_rng(12345)
    p = [[(rand() * 2 - 1) * 20, (rand() * 2 - 1) * 20, (rand() * 2 - 1) * 20] for _ in range(n)]
    k = 14
    for it in range(500):
        t = 3 * (1 - it / 500) + 0.05
        f = [[0, 0, 0] for _ in range(n)]
        for i in range(n):
            for j in range(i + 1, n):
                d = [p[i][0] - p[j][0], p[i][1] - p[j][1], p[i][2] - p[j][2]]
                dist = max(0.01, math.hypot(*d))
                r = k * k / dist
                for a in range(3):
                    f[i][a] += d[a] / dist * r
                    f[j][a] -= d[a] / dist * r
        for i, j, w in edges:
            d = [p[i][0] - p[j][0], p[i][1] - p[j][1], p[i][2] - p[j][2]]
            dist = max(0.01, math.hypot(*d))
            a1 = dist * dist / k * (0.6 + 0.2 * w)
            for a in range(3):
                f[i][a] -= d[a] / dist * a1
                f[j][a] += d[a] / dist * a1
        for i in range(n):
            dist = max(0.01, math.hypot(*f[i]))
            for a in range(3):
                p[i][a] += f[i][a] / dist * min(dist, t)
    c = [sum(q[a] for q in p) / n for a in range(3)]
    for q in p:
        for a in range(3):
            q[a] -= c[a]
    R = max(math.hypot(*q) for q in p)
    sc = 36 / R
    return [[jsnum(to_fixed(v * sc, 2)) for v in q] for q in p]


def utm(lat, lon):
    a = 6378137
    f = 1 / 298.257223563
    k0 = 0.9996
    e2 = f * (2 - f)
    ep2 = e2 / (1 - e2)
    lon0 = -75 * math.pi / 180
    p = lat * math.pi / 180
    lam = lon * math.pi / 180
    N = a / math.sqrt(1 - e2 * math.sin(p) ** 2)
    T = math.tan(p) ** 2
    C = ep2 * math.cos(p) ** 2
    A = math.cos(p) * (lam - lon0)
    M = a * ((1 - e2 / 4 - 3 * e2 * e2 / 64 - 5 * e2 ** 3 / 256) * p
             - (3 * e2 / 8 + 3 * e2 * e2 / 32 + 45 * e2 ** 3 / 1024) * math.sin(2 * p)
             + (15 * e2 * e2 / 256 + 45 * e2 ** 3 / 1024) * math.sin(4 * p)
             - (35 * e2 ** 3 / 3072) * math.sin(6 * p))
    E = 500000 + k0 * N * (A + (1 - T + C) * A ** 3 / 6
                           + (5 - 18 * T + T * T + 72 * C - 58 * ep2) * A ** 5 / 120)
    Nn = k0 * (M + N * math.tan(p) * (A * A / 2
                                      + (5 - T + 9 * C + 4 * C * C) * A ** 4 / 24
                                      + (61 - 58 * T + T * T + 600 * C - 330 * ep2) * A ** 6 / 720))
    return E, Nn


def to_scene(lat, lon):
    E, N = utm(lat, lon)
    return {'x': jsnum(to_fixed(((E - 586865) - 5341.33) / 10, 2)),
            'z': jsnum(to_fixed(-(((N - 509725) - 3161.9) / 10), 2))}


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    def sh(name):
        w = wb[name] if name in wb.sheetnames else None
        if w is None:
            raise RuntimeError('Falta la hoja ' + name)
        return w

    capas = [{'code': r[0], 'name': r[1], 'desc': r[2], 'color': r[3].upper(),
              'icon': r[4], 'badge': r[5]} for r in rows(sh('CAPAS'), 1, 6)]
    tipos = [{'id': i, 'code': r[0], 'name': r[1], 'def': r[2],
              'color': r[3].upper(), 'icon': r[4]} for i, r in enumerate(rows(sh('INTERACCIONES'), 1, 5))]
    capa_idx = {c['name']: i for i, c in enumerate(capas)}
    tipo_idx = {t['name']: t['id'] for t in tipos}
    errs = []

    def num_or_nan(s):
        try:
            return float(s)
        except (ValueError, TypeError):
            return float('nan')

    nodes = []
    for r in rows(sh('NODOS'), 1, 18):
        if r[2] not in capa_idx:
            errs.append('Nodo %s: capa no valida "%s"' % (r[0], r[2]))
        lat = num_or_nan(r[12])
        lon = num_or_nan(r[13])
        if math.isnan(lat) or math.isnan(lon):
            errs.append('Nodo %s: falta latitud/longitud' % r[0])
            geo = {'lat': 0, 'lon': 0, 'x': 0, 'z': 0}
        else:
            geo = {'lat': jsnum(lat), 'lon': jsnum(lon), **to_scene(lat, lon)}
        nodes.append({'id': r[0], 'name': r[1], 'cat': capa_idx.get(r[2]), 'sciname': r[3],
                      'scale': r[4], 'loc': r[5], 'role': r[6], 'alert': r[7], 'actors': r[8],
                      'hypothesis': r[9], 'source': r[10], 'img': r[11], 'img1': r[14],
                      'img2': r[15], 'icon': r[17], 'geo': geo})
    for nd in nodes:
        if not nd['icon'] and nd['cat'] is not None:
            nd['icon'] = capas[nd['cat']]['icon']

    def inside(p):
        return -182.5 <= p['x'] <= 534.1 and abs(p['z']) <= 316.2

    for nd in nodes:
        if nd['geo'] and not (isinstance(nd['geo']['x'], float) and math.isnan(nd['geo']['x'])) \
                and not inside(nd['geo']):
            print('AVISO: fuera del mapa 3D de Kennedy (no se vera en Territorio): '
                  + nd['id'] + ' - ' + nd['name'])

    id_idx = {n['id']: i for i, n in enumerate(nodes)}
    if len(id_idx) != len(nodes):
        errs.append('IDs de nodo duplicados')

    edges = []
    for r in rows(sh('ARISTAS'), 1, 12):
        if r[1] not in id_idx:
            errs.append('Arista %s: origen inexistente "%s"' % (r[0], r[1]))
        if r[2] not in id_idx:
            errs.append('Arista %s: destino inexistente "%s"' % (r[0], r[2]))
        if r[3] not in tipo_idx:
            errs.append('Arista %s: tipo no valido "%s"' % (r[0], r[3]))
        try:
            w = float(r[7])
        except (ValueError, TypeError):
            w = float('nan')
        if w != w or w == 0:
            w = 1
        else:
            w = jsnum(w)
        edges.append({'id': r[0], 'source': id_idx.get(r[1]), 'target': id_idx.get(r[2]),
                      'type': tipo_idx.get(r[3]), 'detail': r[4], 'rationale': r[5],
                      'tension': r[6], 'weight': w, 'declared': r[8], 'operative': r[9],
                      'hypothesis': r[10], 'evidence': r[11]})
    if errs:
        print('ERRORES:\n - ' + '\n - '.join(errs))
        sys.exit(1)

    n = len(nodes)
    adj = [[] for _ in range(n)]
    we = []
    for e in edges:
        adj[e['source']].append(e['target'])
        adj[e['target']].append(e['source'])
        we.append([e['source'], e['target'], e['weight']])

    eff_name = ''
    if 'RED' in wb.sheetnames:
        for k, v in rows(wb['RED'], 1, 2):
            if k == 'capa_efectos':
                eff_name = v
    eff_idx = next((i for i, c in enumerate(capas) if c['name'] == eff_name), -1) if eff_name else -1

    bc = brandes(n, adj)
    comm, Q = louvain(n, we)

    vis = [i for i in range(n) if nodes[i]['cat'] != eff_idx]
    vmap = {g: k for k, g in enumerate(vis)}
    wsub = [[vmap[a], vmap[b], w] for a, b, w in we if a in vmap and b in vmap]
    psub = layout3d(len(vis), wsub) if eff_idx >= 0 else layout3d(n, we)
    if eff_idx >= 0:
        pos = []
        for i in range(n):
            if i in vmap:
                pos.append(psub[vmap[i]])
            else:
                nb = [j for j in adj[i] if j in vmap]
                if not nb:
                    pos.append([0, 0, 0])
                else:
                    m = [sum(psub[vmap[j]][kk] for j in nb) / len(nb) for kk in range(3)]
                    r2 = make_rng(i * 7919 + 13)
                    pos.append([jsnum(to_fixed(v + (r2() - 0.5) * 6, 2)) for v in m])
    else:
        pos = psub

    gephi = {'Q': Q, 'nodes': {}}
    for i, nd in enumerate(nodes):
        gephi['nodes'][nd['id']] = [comm[i], bc[i], pos[i][0], pos[i][1], pos[i][2]]

    config = {'faq': []}
    if 'RED' in wb.sheetnames:
        for k, v in rows(wb['RED'], 1, 2):
            m = re.match(r'^faq(\d+)_(pregunta|respuesta)$', k)
            if m:
                i = int(m.group(1)) - 1
                while len(config['faq']) <= i:
                    config['faq'].append({'q': '', 'a': ''})
                config['faq'][i]['q' if m.group(2) == 'pregunta' else 'a'] = v
            else:
                config[k] = v
        config['faq'] = [fl for fl in config['faq'] if fl and fl['q']]

    data = {'config': config, 'categories': capas, 'interactions': tipos,
            'nodes': nodes, 'edges': edges, 'gephi': gephi}
    with open(HTML_PATH, 'r', encoding='utf-8', newline='') as f:
        html = f.read()
    A = '<!--RED_DATA_START-->'
    B = '<!--RED_DATA_END-->'
    ia = html.find(A)
    ib = html.find(B)
    if ia < 0 or ib < 0:
        raise RuntimeError('La pagina no tiene los marcadores %s ... %s' % (A, B))
    payload = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('<', '\\u003c')
    html = html[:ia] + A + '<script>window.RED_DATA = ' + payload + ';</script>' + html[ib:]
    with open(HTML_PATH, 'w', encoding='utf-8', newline='') as f:
        f.write(html)
    print('OK: %d nodos, %d aristas, %d capas, %d tipos, %d comunidades (Q=%s) -> %s'
          % (n, len(edges), len(capas), len(tipos), len(set(comm)), Q, HTML_PATH.name))


if __name__ == '__main__':
    main()
